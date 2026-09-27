"""Tabular input decoders shared by Workflow CLI execution.

Supported file types are detected from the extension: csv, tsv, jsonl, json,
xlsx. This module only parses provided bytes/text; sandbox file I/O belongs to CLI.
"""
from __future__ import annotations

import csv
import io
import json

from vibecanvas_api.agents.tools.decorator import ToolError

# ── read ──────────────────────────────────────────────────────────────────

def _resolve_sheet(path: str, names: list[str], sheet: str) -> str:
    """Pick the target sheet, or raise a CLEAN ToolError. ``sheet=''`` + a single
    sheet → that sheet; ``''`` + multiple sheets → ``sheet_required`` (lists names);
    a named sheet must exist → else ``sheet_not_found`` (lists names)."""
    sheet = (sheet or "").strip()
    if sheet:
        if sheet in names:
            return sheet
        raise ToolError("sheet_not_found", f"sheet {sheet!r} not found in {path!r}; "
                        f"available sheets: {', '.join(names)}")
    if len(names) == 1:
        return names[0]
    raise ToolError("sheet_required", f"path {path!r} has multiple sheets "
                    f"({', '.join(names)}); specify one with the sheet argument")


def _xlsx_to_rows(data: bytes, path: str, sheet: str) -> tuple[list[dict], list[str]]:
    """Return (rows, columns). ``columns`` is the authoritative header (non-empty
    cells of the first row) — preserved even when there are zero data rows."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ToolError("read_failed", f"could not read {path!r} as a spreadsheet")
    try:
        chosen = _resolve_sheet(path, wb.sheetnames, sheet)
        it = wb[chosen].iter_rows(values_only=True)
        try:
            headers = ["" if h is None else str(h) for h in next(it)]
        except StopIteration:
            return [], []
        cols = [h for h in headers if h]
        out = []
        for vals in it:
            out.append({headers[i]: (vals[i] if i < len(vals) else None)
                        for i in range(len(headers)) if headers[i]})
        return out, cols
    except ToolError:
        raise
    except Exception:
        raise ToolError("read_failed", f"could not read {path!r} as a spreadsheet")
    finally:
        wb.close()


def _text_to_rows(text: str, ext: str) -> tuple[list[dict], list[str]] | None:
    """Return (rows, columns) or None for an unrecognized format. For csv/tsv the
    columns are the header (kept even with zero data rows); for jsonl/json there is no
    separate header, so columns are the union of the rows' keys."""
    if ext in ("csv", "tsv"):
        delim = "\t" if ext == "tsv" else ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delim)
        rows = list(reader)
        return rows, [c for c in (reader.fieldnames or []) if c is not None]
    if ext == "jsonl":
        rows = [json.loads(ln) for ln in text.splitlines() if ln.strip()]
        return rows, _all_columns(rows)
    if ext == "json":
        obj = json.loads(text or "[]")
        if isinstance(obj, dict) and "rows" in obj:
            rows = obj["rows"]
        elif isinstance(obj, list):
            rows = obj
        else:
            return None
        return rows, _all_columns(rows)
    return None


def _all_columns(rows: list[dict]) -> list[str]:
    cols: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in cols:
                cols.append(k)
    return cols
